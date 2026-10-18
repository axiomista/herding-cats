"""Builds the format test corpus in seed/formats/ and its section of seed/answer_key.json.

Each file holds known commitments (C10 to C32) in a different format, so a live run shows what
the cycle does with each format: scanned, skipped, or failed, and which commitments it finds.
Every evidence string is defined once here and written into both the file and the answer key,
so the two cannot drift apart. The files are committed; rerun this only to change them.

Run from the repo root: python3 seed/formats/build.py
Needs python-docx, python-pptx, openpyxl, reportlab, Pillow, qpdf, ffmpeg and Noto Sans CJK.
"""
import json, os, random, re, subprocess, zipfile
from datetime import date

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Image as RLImage, Paragraph, SimpleDocTemplate, Spacer, Table

OUT = os.path.dirname(os.path.abspath(__file__))
ANSWER_KEY = os.path.join(OUT, "..", "answer_key.json")
random.seed(7)

SANS = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
SANS_B = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
SERIF = "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf"
CJK = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"

key = {"files": {}, "handling": {}, "commitments": [], "non_commitments": []}

# What the code does today with each file (src/engine/cycle.ts SUPPORTED and maxFileBytes, DECISIONS #18).
HANDLING = {
    "scanned": "Supported type under 5 MB: classified, extracted and evidence-checked.",
    "scanned_no_text_layer": "PDF with no text layer. Scanned (or an error if Box has no extracted text for it), "
                             "but /ai/extract has no OCR and there is no text to check evidence against, so its "
                             "commitments are not expected to be found yet.",
    "skipped_type": "Extension not in the supported list: skipped by the inventory, no AI call.",
    "skipped_size": "Over the 5 MB limit: skipped by the inventory, no AI call.",
    "error": "Scanned, but Box cannot produce its text: expect an entry in the page's errors, and the cycle continues.",
}


def path(name):
    return os.path.join(OUT, name)


def register(fid, name, fmt, handling, note=""):
    assert handling in HANDLING, handling
    key["files"][fid] = name
    key["handling"][fid] = {"format": fmt, "expected": handling, "note": note}


def commit(cid, fid, owner, due, action, evidence, expected="untracked", **extra):
    key["commitments"].append(dict(id=cid, source=fid, owner=owner, due=due, action=action,
                                   evidence=evidence, expected_finding=expected, **extra))
    return evidence


def non(fid, kind, text):
    key["non_commitments"].append(dict(source=fid, kind=kind, text=text))
    return text


NOT_SCANNED = "not found: file skipped"
NO_OCR = "not found until an OCR path exists; then untracked at low certainty (evidence unverifiable)"

# ---------- image helpers ----------


def font(p, size, index=0):
    return ImageFont.truetype(p, size, index=index)


def wrap(draw, text, f, width):
    words, lines, cur = text.split(" "), [], ""
    for w in words:
        trial = (cur + " " + w).strip()
        if draw.textlength(trial, font=f) <= width:
            cur = trial
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def scan_page(blocks, size=(1275, 1650), margin=110):
    """Render (text, style) blocks to a letter page at 150 dpi, then age it like a scan."""
    img = Image.new("L", size, 250)
    d = ImageDraw.Draw(img)
    y = margin
    for text, style in blocks:
        if style == "gap":
            y += 28
            continue
        f = {"h1": font(SERIF, 40), "h2": font(SANS_B, 28), "body": font(SERIF, 26),
             "small": font(SERIF, 22), "sig": font(SERIF, 30)}[style]
        for line in wrap(d, text, f, size[0] - 2 * margin):
            d.text((margin, y), line, font=f, fill=25)
            y += int(f.size * 1.45)
        y += 8
    if any(s == "sig" for _, s in blocks):  # a scribbled signature under the last sig line
        pts = [(margin + 20 + i * 14, y + 30 + random.randint(-18, 18)) for i in range(22)]
        d.line(pts, fill=40, width=3)
    img = img.rotate(random.uniform(-0.9, 0.9), resample=Image.BICUBIC, fillcolor=245)
    noise = Image.effect_noise(size, 18).point(lambda v: 0 if v < 120 else 255)
    img = Image.blend(img, Image.composite(img, noise, noise.point(lambda v: 255 if v > 30 else 0)), 0.15)
    return img.filter(ImageFilter.GaussianBlur(0.6)).convert("RGB")


def save_image_pdf(pages, name):
    pages[0].save(path(name), "PDF", resolution=150, save_all=True, append_images=pages[1:])


ss = getSampleStyleSheet()

# =====================================================================
# 1. DOCX: native text, smart quotes and em dashes
# =====================================================================
fid, name = "standup_docx", "2026-10-14 Floor 3 Annexation Standup.docx"
register(fid, name, "docx", "scanned", "Smart quotes, an em dash and bullets test evidence normalization.")
c10 = commit("C10", fid, "Tuxedo Meowtani", "2026-10-30",
             "Finalize the seating chart for the annexation of the Third Floor break room",
             "Tuxedo will finalize the seating chart for the “peaceful annexation” of the Third Floor break room by October 30.")
c11 = commit("C11", fid, "Dale Henderson", "2026-10-27",
             "File form HR-12b (Request for Additional Lap Time)",
             "Dale will file form HR-12b — “Request for Additional Lap Time” — by October 27.")
n1 = non(fid, "idea", "Someone should probably look into the stapler situation.")
n2 = non(fid, "decision", "The Third Floor break room is now Sovereign Executive Territory.")
doc = Document()
doc.add_heading("Floor 3 Annexation — Standup", 1)
doc.add_paragraph("Date: Wednesday, October 14, 2026\nAttendees: Tuxedo Meowtani, Dale Henderson, Pixel\n"
                  "Absent: Sir Reginald Biscuits (in a box; did not respond to knocking)\nNote-taker: Pixel")
doc.add_heading("Updates", 2)
doc.add_paragraph("Floor 3 break-room occupancy by executives has reached 100% between 2 and 4 p.m. "
                  "Associates have been “invited” to use the Floor 2 kitchen instead.", style="List Bullet")
doc.add_paragraph(n2, style="List Bullet")
doc.add_heading("Action items", 2)
doc.add_paragraph(c10, style="List Bullet")
doc.add_paragraph(c11, style="List Bullet")
doc.add_heading("Parking lot", 2)
doc.add_paragraph(n1, style="List Bullet")
doc.save(path(name))

# =====================================================================
# 2. PDF with a text layer and a table
# =====================================================================
fid, name = "stapler_pdf", "2026-10-08 Memo - Stapler Consolidation Initiative.pdf"
register(fid, name, "pdf (text layer)", "scanned", "Table cells must not scramble the evidence sentence.")
c12 = commit("C12", fid, "Pixel", "2026-11-03", "Deliver the Stapler Census to the Chairmeow's office",
             "Pixel will deliver the completed Stapler Census to the Chairmeow’s office by November 3.")
n3 = non(fid, "decision", "Effective immediately, all staplers are the property of the Board.")
n4 = non(fid, "policy", "Associates may request a stapler in writing, in triplicate.")
SimpleDocTemplate(path(name), pagesize=letter).build([
    Paragraph("Nyanland-Meowtani Corporation — Internal Memorandum", ss["Title"]),
    Paragraph("<b>To:</b> All Associates &nbsp; <b>From:</b> Office of the Chairmeow &nbsp; <b>Date:</b> October 8, 2026", ss["Normal"]),
    Paragraph("<b>Re:</b> Stapler Consolidation Initiative", ss["Normal"]), Spacer(1, 14),
    Paragraph("An audit of desk surfaces found that staplers occupy space better used for sitting. " + n3 + " " + n4,
              ss["BodyText"]), Spacer(1, 10),
    Table([["Floor", "Staplers found", "Staplers knocked off desk"], ["2", "41", "39"], ["3", "57", "57"],
           ["4", "12", "pending"]],
          style=[("GRID", (0, 0), (-1, -1), 0.5, "#888888"), ("BACKGROUND", (0, 0), (-1, 0), "#eeeeee")]),
    Spacer(1, 12), Paragraph(c12, ss["BodyText"]), Spacer(1, 10),
    Paragraph("Thank you for your continued compliance.", ss["BodyText"])])

# =====================================================================
# 3. PPTX: one commitment on a slide, one only in speaker notes
# =====================================================================
fid, name = "roadmap_pptx", "Q4 World Domination Roadmap.pptx"
register(fid, name, "pptx", "scanned", "C14 is only in the speaker notes: the run shows whether Box's text includes notes.")
c13 = commit("C13", fid, "Sir Reginald Biscuits", "2026-11-12", "Present the Global Nap Schedule to the Board",
             "Sir Reginald Biscuits will present the Global Nap Schedule to the Board by November 12.")
c14 = commit("C14", fid, "Dale Henderson", "2026-11-10", "Book the projector for the Board meeting",
             "Dale will book the projector for the Board meeting by November 10.",
             expected="untracked if Box's text includes speaker notes; otherwise not found")
n5 = non(fid, "idea", "Explore a subscription tier for premium lap access.")
prs = Presentation()
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = "Q4 World Domination Roadmap"
s.placeholders[1].text = "Nyanland-Meowtani Corporation · Strategy Office · October 2026"
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = "Three Pillars of Gentle Domination"
s.placeholders[1].text_frame.text = "Sunbeam sovereignty"
for t in ["Lap acquisition", "Cardboard infrastructure", n5]:
    s.placeholders[1].text_frame.add_paragraph().text = t
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = "Next Steps"
s.placeholders[1].text_frame.text = c13
s.notes_slide.notes_text_frame.text = "Reminder for the presenter: " + c14
prs.save(path(name))

# =====================================================================
# 4. XLSX: tracker rows, not sentences
# =====================================================================
fid, name = "tracker_xlsx", "Action Item Tracker - Floor 4 Annexation.xlsx"
register(fid, name, "xlsx", "scanned",
         "Evidence is a row, not a sentence, so the evidence check may drop it. "
         "The cancelled row and the Ideas sheet must not become findings.")
rows = [("Tuxedo Meowtani", "Relocate the Floor 4 ficus to a sunnier, more strategic position", date(2026, 10, 29), "Open"),
        ("Dale Henderson", "Label every chair on Floor 4 'Reserved for Management'", date(2026, 11, 5), "Open"),
        ("Pixel", "Requisition three bags of premium catnip", date(2026, 10, 1), "Cancelled")]
for cid, r in (("C15", rows[0]), ("C16", rows[1])):
    commit(cid, fid, r[0], r[2].isoformat(), r[1], r[1],
           expected="untracked if the evidence check accepts the cell text; otherwise unverified", evidence_kind="cell")
non(fid, "cancelled_commitment", rows[2][1])
wb = Workbook()
ws = wb.active
ws.title = "Tracker"
ws.append(["Owner", "Action", "Due", "Status"])
for c in ws[1]:
    c.font = Font(bold=True)
for r in rows:
    ws.append(list(r))
for col, w in zip("ABCD", (22, 62, 12, 11)):
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
    row[0].number_format = "yyyy-mm-dd"
ideas = wb.create_sheet("Ideas")
ideas.append(["Idea", "Champion"])
ideas.append([non(fid, "idea", "Rename Floor 4 'The Annex' for brand cohesion?"), "nobody yet"])
ideas.append([non(fid, "idea", "Mandatory purring breaks (2 per hour)"), "unclear"])
wb.save(path(name))

# =====================================================================
# 5. Scanned PDF, image only, 1 page
# =====================================================================
fid, name = "decree_scan", "2026-10-09 Decree - Mandatory Sunbeam Rotation (signed scan).pdf"
register(fid, name, "pdf (image only)", "scanned_no_text_layer")
c17 = commit("C17", fid, "Pixel", "2026-10-31", "Post the Sunbeam Rotation Schedule on every floor",
             "Pixel will post the new Sunbeam Rotation Schedule on every floor by October 31.", expected=NO_OCR)
n6 = non(fid, "decision", "By decree of the Chairmeow, sunbeams shall rotate among executives hourly.")
save_image_pdf([scan_page([
    ("OFFICIAL DECREE No. 0042", "h1"), ("Office of the Chairmeow · October 9, 2026", "small"), ("", "gap"),
    ("Subject: Mandatory Sunbeam Rotation", "h2"), ("", "gap"),
    ("Recent disputes over the window ledge on Floor 2 have disrupted three naps and one quarterly review.", "body"),
    (n6, "body"), ("Associates are reminded that sunbeams are not a human entitlement.", "body"), ("", "gap"),
    (c17, "body"), ("", "gap"), ("", "gap"),
    ("Signed, Chairmeow Baroness Nyanland", "sig")])], name)

# =====================================================================
# 6. Scanned PDF, 7 pages: Box AI reads at most 5 images or pages per request
# =====================================================================
fid, name = "reserves_scan7", "FY26 Annual Report - Treat Reserves (scanned, 7 pages).pdf"
register(fid, name, "pdf (image only, 7 pages)", "scanned_no_text_layer",
         "C19 is on page 2 and C18 on page 6. With an OCR path, a 5-page limit would drop C18 silently.")
c19 = commit("C19", fid, "Tuxedo Meowtani", "2026-11-04", "Audit the catnip vault",
             "Tuxedo will audit the catnip vault by November 4.", expected=NO_OCR, page=2)
c18 = commit("C18", fid, "Sir Reginald Biscuits", "2026-11-20", "Reconcile the salmon reserve ledger",
             "Sir Reginald Biscuits will reconcile the salmon reserve ledger by November 20.",
             expected=NO_OCR + "; missing even then if only 5 pages are read", page=6)
filler = ("Reserves remained stable throughout the period. Variance is attributed to seasonal "
          "snacking and one unexplained incident in the pantry. No further comment will be made.")
pages = []
for i in range(1, 8):
    blocks = [(f"FY26 Annual Report — Treat Reserves · Page {i} of 7", "small"), ("", "gap"),
              (["Executive Summary", "Catnip Holdings", "Tuna Futures", "Kibble Liquidity",
                "Cardboard Assets", "Salmon Reserves", "Outlook"][i - 1], "h2"), ("", "gap"), (filler, "body"),
              (f"Table {i}.1: Units held at quarter end: {random.randint(200, 900)}; "
               f"units consumed by executives: {random.randint(200, 900)}.", "body")]
    if i == 2:
        blocks += [("", "gap"), (c19, "body")]
    if i == 6:
        blocks += [("", "gap"), (c18, "body")]
    pages.append(scan_page(blocks))
save_image_pdf(pages, name)

# =====================================================================
# 7. Mixed PDF: two text pages and a scanned signature page
# =====================================================================
fid, name = "contract_mixed", "2026-10-10 Contract - Laser Pointer Maintenance (with signed page).pdf"
register(fid, name, "pdf (pages 1-2 text, page 3 scanned)", "scanned",
         "C21 is on a text page; C20 is only on the scanned page.")
c21 = commit("C21", fid, "Pixel", "2026-11-02", "Schedule the quarterly red-dot calibration",
             "Pixel will schedule the quarterly red-dot calibration with Laser Dynamics by November 2.")
c20 = commit("C20", fid, "Dale Henderson", "2026-10-28", "Return the countersigned contract to Laser Dynamics",
             "Dale will return the countersigned contract to Laser Dynamics by October 28.", expected=NO_OCR)
n7 = non(fid, "external_request", "Laser Dynamics requires 48 hours' notice before any calibration visit.")
tmp_text, tmp_scan = path("_tmp_text.pdf"), "_tmp_scan.pdf"
story = [Paragraph("Service Agreement — Laser Pointer Maintenance", ss["Title"]),
         Paragraph("Between Nyanland-Meowtani Corporation and Laser Dynamics LLC · October 10, 2026", ss["Normal"]),
         Spacer(1, 12)]
for clause in range(1, 9):
    story += [Paragraph(f"<b>Clause {clause}.</b> The red dot shall remain catchable-adjacent but never caught. "
                        "Calibration shall preserve the dot's essential elusiveness.", ss["BodyText"]), Spacer(1, 6)]
story += [Paragraph(n7, ss["BodyText"]), Spacer(1, 400), Paragraph(c21, ss["BodyText"])]
SimpleDocTemplate(tmp_text, pagesize=letter).build(story)
save_image_pdf([scan_page([("SIGNATURE PAGE", "h1"), ("", "gap"),
                           ("Executed by the parties on the date below.", "body"), ("", "gap"),
                           (c20, "body"), ("", "gap"),
                           ("For Laser Dynamics LLC: Gary Okafor", "sig")])], tmp_scan)
subprocess.run(["qpdf", "--empty", "--pages", tmp_text, "1-2", path(tmp_scan), "--", path(name)], check=True)
os.remove(tmp_text)
os.remove(path(tmp_scan))

# =====================================================================
# 8. PNG whiteboard: shorthand owner and date
# =====================================================================
fid, name = "whiteboard_png", "Whiteboard - Operation Lap Takeover.png"
register(fid, name, "png", "skipped_type", "OCR candidate. Shorthand owner 'TUX' and date 'Nov 7'.")
wb_line = "TUX: order 200 heated lap blankets by Nov 7"
commit("C22", fid, "Tuxedo Meowtani", "2026-11-07", "Order 200 heated lap blankets", wb_line, expected=NOT_SCANNED)
img = Image.new("RGB", (1024, 768), (246, 247, 244))
d = ImageDraw.Draw(img)
d.rectangle([12, 12, 1011, 755], outline=(170, 170, 170), width=6)
d.text((60, 50), "OPERATION LAP TAKEOVER", font=font(SANS_B, 54), fill=(30, 60, 160))
d.text((60, 130), "Phase 1: warm the laps   Phase 2: occupy the laps", font=font(SANS, 30), fill=(40, 40, 40))
d.text((60, 230), "IDEAS", font=font(SANS_B, 36), fill=(170, 30, 30))
for i, t in enumerate([non(fid, "idea", "lap-sharing app??"), non(fid, "idea", "humans sit still = reward")]):
    d.text((80, 285 + i * 50), "- " + t, font=font(SANS, 32), fill=(40, 40, 40))
d.text((60, 420), "ACTIONS", font=font(SANS_B, 36), fill=(20, 120, 50))
d.text((80, 475), "- " + wb_line, font=font(SANS, 34), fill=(20, 20, 20))
d.ellipse([700, 560, 960, 720], outline=(170, 30, 30), width=5)
d.text((735, 615), "NO DOGS", font=font(SANS_B, 38), fill=(170, 30, 30))
img.rotate(-1.2, resample=Image.BICUBIC, fillcolor=(230, 230, 228)).filter(ImageFilter.GaussianBlur(0.7)).save(path(name))

# =====================================================================
# 9. TIFF notice: recurring commitment, no date
# =====================================================================
fid, name = "notice_tiff", "Break Room Notice - Treat Dispenser.tiff"
register(fid, name, "tiff", "skipped_type", "OCR candidate. Recurring commitment with no due date.")
c23 = commit("C23", fid, "Dale Henderson", None, "Refill the executive treat dispenser every Friday",
             "Dale will refill the executive treat dispenser every Friday.", expected=NOT_SCANNED)
scan_page([("NOTICE", "h1"), ("", "gap"), ("Break Room, Floor 3", "h2"), ("", "gap"),
           ("The executive treat dispenser is for executives.", "body"),
           (c23, "body"), ("Associates found near the dispenser will be gently stared at.", "body")],
          size=(1024, 900), margin=70).save(path(name), compression="tiff_lzw")

# =====================================================================
# 10. WEBP chat screenshot: the owner is the message author
# =====================================================================
fid, name = "chat_webp", "Chat Screenshot - #exec-lounge.webp"
register(fid, name, "webp", "skipped_type", "OCR candidate. 'I' resolves to the message author (Pixel).")
chat = [("Tuxedo Meowtani", "10:02", "who is getting the nap pod quote"),
        ("Pixel", "10:04", "I'll get the Nap Pod quote from Snoozeworks by Oct 30"),
        ("Sir Reginald Biscuits", "10:05", "excellent. pods for executives only"),
        ("Dale Henderson", "10:09", "can associates use them on weekends?"),
        ("Sir Reginald Biscuits", "10:10", "no")]
commit("C24", fid, "Pixel", "2026-10-30", "Get the Nap Pod quote from Snoozeworks", chat[1][2], expected=NOT_SCANNED)
non(fid, "open_question", chat[3][2])
img = Image.new("RGB", (900, 620), (255, 255, 255))
d = ImageDraw.Draw(img)
d.rectangle([0, 0, 900, 60], fill=(74, 21, 75))
d.text((24, 14), "# exec-lounge", font=font(SANS_B, 26), fill=(255, 255, 255))
y = 90
for who, t, msg in chat:
    d.rounded_rectangle([24, y, 70, y + 46], 8, fill=(random.randint(80, 200), 120, 160))
    d.text((88, y - 2), who, font=font(SANS_B, 22), fill=(20, 20, 20))
    d.text((88 + d.textlength(who, font=font(SANS_B, 22)) + 12, y + 2), t, font=font(SANS, 17), fill=(120, 120, 120))
    d.text((88, y + 28), msg, font=font(SANS, 22), fill=(30, 30, 30))
    y += 100
img.save(path(name), "WEBP", quality=90)

# =====================================================================
# 11. Japanese DOCX
# =====================================================================
fid, name = "tokyo_docx", "2026-10-13 東京支社 定例会議 議事録.docx"
register(fid, name, "docx (Japanese)", "scanned",
         "Katakana aliases resolve (タキシード = Tuxedo, ピクセル = Pixel); タマ支社長 has no Box user and is flagged; "
         "来週金曜日 ('next Friday') is relative and is flagged, not guessed.")
c25 = commit("C25", fid, "Tuxedo Meowtani", "2026-11-14", "Draft the Tokyo branch nap schedule",
             "タキシードは、11月14日までに東京支社の昼寝スケジュール案を作成する。",
             translation="Tuxedo will draft the Tokyo branch nap schedule by November 14.")
c26 = commit("C26", fid, "タマ支社長", "2026-10-31", "Survey all employees' lap usage",
             "タマ支社長は、10月31日までに全社員の膝の使用状況を調査する。",
             expected="untracked at low certainty (owner unresolved)",
             translation="Branch Manager Tama will survey all employees' lap usage by October 31.")
c27 = commit("C27", fid, "Pixel", None, "Update the sunbeam map",
             "ピクセルは来週金曜日までに日なたマップを更新する。",
             expected="untracked at low certainty (relative date flagged)",
             translation="Pixel will update the sunbeam map by next Friday.")
n8 = non(fid, "open_question", "人間にも昼寝を許可すべきか？（継続審議）")
n9 = non(fid, "decision", "段ボール箱は役員専用とする。")
doc = Document()
st = doc.styles["Normal"]
st.font.name = "Noto Sans CJK JP"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK JP")
st.font.size = Pt(11)
doc.add_heading("ニャンランド・ニャウタニ株式会社　東京支社 定例会議 議事録", 1)
doc.add_paragraph("日時：2026年10月13日（火）15:00〜15:20（昼寝のため短縮）\n"
                  "出席者：タキシード・ニャウタニ、タマ支社長、ピクセル\n欠席：ビスケッツ卿（箱の中）\n書記：ピクセル")
doc.add_heading("報告事項", 2)
doc.add_paragraph("東京支社の日当たりは本社より12％良好である。窓際の席はすでに役員が占拠済み。", style="List Bullet")
doc.add_paragraph(n9, style="List Bullet")
doc.add_heading("決定事項・宿題", 2)
for c in (c25, c26, c27):
    doc.add_paragraph(c, style="List Bullet")
doc.add_heading("継続審議", 2)
doc.add_paragraph(n8, style="List Bullet")
doc.save(path(name))

# =====================================================================
# 12. Japanese PNG whiteboard: Japanese OCR and a Reiwa-era date
# =====================================================================
fid, name = "tokyo_png", "ホワイトボード - 世界征服会議.png"
register(fid, name, "png (Japanese)", "skipped_type",
         "OCR candidate. Date written in the Reiwa era (令和8年 = 2026); ビスケッツ卿 = Sir Reginald Biscuits.")
jp_line = "ビスケッツ卿：令和8年11月20日までに缶切り研修第2章を完成"
commit("C28", fid, "Sir Reginald Biscuits", "2026-11-20", "Finish chapter 2 of the can-opening training", jp_line,
       expected=NOT_SCANNED,
       translation="Sir Reginald Biscuits: finish chapter 2 of the can-opening training by November 20, Reiwa 8 (2026).")
img = Image.new("RGB", (1024, 768), (247, 247, 243))
d = ImageDraw.Draw(img)
d.rectangle([12, 12, 1011, 755], outline=(160, 160, 160), width=6)
d.text((50, 45), "世界征服会議（第7回）", font=font(CJK, 56), fill=(30, 60, 160))
d.text((50, 130), "令和8年10月14日　議題：人間の膝の有効活用", font=font(CJK, 30), fill=(50, 50, 50))
d.text((50, 220), "アイデア", font=font(CJK, 38), fill=(170, 30, 30))
for i, t in enumerate([non(fid, "idea", "膝のサブスク？"), non(fid, "idea", "人間用キャットタワー")]):
    d.text((70, 280 + i * 52), "・" + t, font=font(CJK, 34), fill=(40, 40, 40))
d.text((50, 430), "宿題", font=font(CJK, 38), fill=(20, 120, 50))
d.text((70, 490), "・" + jp_line, font=font(CJK, 28), fill=(20, 20, 20))
img.rotate(0.8, resample=Image.BICUBIC, fillcolor=(232, 232, 228)).filter(ImageFilter.GaussianBlur(0.6)).save(path(name))

# =====================================================================
# 13. VTT meeting transcript: plain text with speaker labels
# =====================================================================
fid, name = "rehearsal_vtt", "2026-10-14 All-Paws Rehearsal - Transcript.vtt"
register(fid, name, "vtt", "skipped_type",
         "Plain text, but .vtt is not in the supported list. If it is added, 'I' should resolve to the speaker label.")
c29_text = "I'll print the 400 agendas by October 22."
commit("C29", fid, "Dale Henderson", "2026-10-22", "Print the 400 All-Paws agendas", c29_text, expected=NOT_SCANNED)
lines = [("Tuxedo Meowtani", "Okay, from the top. The Great Hall, the radiator, the podium."),
         ("Tuxedo Meowtani", non(fid, "idea", "We could maybe add a laser show at the end.")),
         ("Pixel", "Do we have agendas?"),
         ("Dale Henderson", "Not yet. " + c29_text),
         ("Sir Reginald Biscuits", "Make them smaller. Humans read too slowly anyway.")]
with open(path(name), "w") as f:
    f.write("WEBVTT\n\n")
    for i, (who, txt) in enumerate(lines):
        f.write(f"{i + 1}\n00:00:{i * 6:02d}.000 --> 00:00:{i * 6 + 5:02d}.500\n<v {who}>{txt}\n\n")

# =====================================================================
# 14 to 19. Files that should be skipped or fail cleanly
# =====================================================================
name = "2026-10-15 Board Meeting Recording.mp4"
register("recording_mp4", name, "mp4", "skipped_type")
subprocess.run(["ffmpeg", "-loglevel", "error", "-y", "-f", "lavfi", "-i", "color=c=0x4a154b:s=640x360:d=5",
                "-f", "lavfi", "-i", "sine=frequency=440:duration=5", "-vf",
                f"drawtext=fontfile={SANS_B}:text='BOARD MEETING (all asleep)':fontcolor=white:fontsize=32:x=(w-tw)/2:y=(h-th)/2",
                "-shortest", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", "-fflags", "+bitexact",
                path(name)], check=True)

name = "Archive - Old Decrees.zip"
register("archive_zip", name, "zip", "skipped_type", "Not unpacked.")
with zipfile.ZipFile(path(name), "w") as z:
    z.writestr(zipfile.ZipInfo("Decree 0001.txt", (2026, 1, 5, 9, 0, 0)),
               "By decree of the Chairmeow, Mondays are abolished.\n")

name = "Floor Plan - Ultra High Resolution (oversize).pdf"
register("oversize_pdf", name, "pdf over 5 MB", "skipped_size", "Holds C30, so a failed size check shows up as a finding.")
c30 = commit("C30", "oversize_pdf", "Pixel", "2026-11-01", "Measure every windowsill on Floor 5",
             "Pixel will measure every windowsill on Floor 5 by November 1.", expected=NOT_SCANNED)
noise = Image.frombytes("RGB", (1350, 1350), random.randbytes(1350 * 1350 * 3))
noise.save(path("_noise.png"))
SimpleDocTemplate(path(name), pagesize=letter).build(
    [Paragraph("Floor 5 Plan (ultra high resolution)", ss["Title"]), Paragraph(c30, ss["BodyText"]),
     RLImage(path("_noise.png"), width=6.5 * inch, height=6.5 * inch)])
os.remove(path("_noise.png"))
assert os.path.getsize(path(name)) > 5 * 1024 * 1024, "the oversize PDF must exceed the 5 MB limit"

name = "Confidential - Succession Plan (password tuna).pdf"
register("encrypted_pdf", name, "pdf, password-protected (password: tuna)", "error", "Holds C31.")
c31 = commit("C31", "encrypted_pdf", "Tuxedo Meowtani", "2026-12-01", "Groom a successor for the Chairmeow",
             "Tuxedo will groom a suitable successor for the Chairmeow by December 1.", expected="not found: unreadable")
SimpleDocTemplate(path("_plain.pdf"), pagesize=letter).build(
    [Paragraph("Succession Plan — Eyes Only", ss["Title"]), Paragraph(c31, ss["BodyText"])])
subprocess.run(["qpdf", "--encrypt", "tuna", "tuna", "256", "--", path("_plain.pdf"), path(name)], check=True)
os.remove(path("_plain.pdf"))

name = "MEMO_FROM_THE_CHAIRMEOW"
register("no_extension", name, "text, no file extension", "skipped_type", "Holds C32.")
c32 = commit("C32", "no_extension", "Sir Reginald Biscuits", "2026-11-15", "Polish the throne",
             "Sir Reginald Biscuits will polish the throne by November 15.", expected=NOT_SCANNED)
with open(path(name), "w") as f:
    f.write("MEMO\nFrom: The Chairmeow\n\n" + c32 + "\n")

name = "Empty Agenda.txt"
register("empty_txt", name, "txt, 0 bytes", "scanned", "No findings; an error entry from Box AI is acceptable, a crash is not.")
open(path(name), "w").close()

# ---------- answer key: replace the "formats" section of seed/answer_key.json ----------
section = {
    "_about": "Format test corpus in seed/formats/, generated by seed/formats/build.py. Upload these files to "
              "The Corporation to see what the cycle does with each format. They do not change C01 to C09; "
              "npm run eval reports them separately and they never fail it.",
    "handling_legend": HANDLING,
    **key,
}
text = open(ANSWER_KEY, encoding="utf-8").read()
cut = text.find(',\n  "formats": ')
body = (text[:cut] if cut >= 0 else text.rstrip().rstrip("}").rstrip()) + ",\n"
block = json.dumps(section, ensure_ascii=False, indent=2).replace("\n", "\n  ")
open(ANSWER_KEY, "w", encoding="utf-8").write(f'{body}  "formats": {block}\n}}\n')
json.loads(open(ANSWER_KEY, encoding="utf-8").read())  # still valid JSON
print(f"built {len(key['files'])} files, {len(key['commitments'])} commitments, "
      f"{len(key['non_commitments'])} non-commitments")
